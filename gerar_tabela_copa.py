"""
Script para gerar a tabela de jogos da Copa do Mundo FIFA 2026
em formato Excel (.xlsx)
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Dados extraídos do site da FIFA (horários no fuso de Brasília)
jogos = [
    # Rodada 1
    ("Grupo A", 1, "México", "África do Sul", "11/06/2026", "16:00"),
    ("Grupo A", 1, "Coreia do Sul", "República Tcheca", "11/06/2026", "23:00"),
    ("Grupo B", 1, "Canadá", "Bósnia e Herzegovina", "12/06/2026", "16:00"),
    ("Grupo D", 1, "EUA", "Paraguai", "12/06/2026", "22:00"),
    ("Grupo B", 1, "Catar", "Suíça", "13/06/2026", "16:00"),
    ("Grupo C", 1, "Brasil", "Marrocos", "13/06/2026", "19:00"),
    ("Grupo C", 1, "Haiti", "Escócia", "13/06/2026", "22:00"),
    ("Grupo D", 1, "Austrália", "Turquia", "14/06/2026", "01:00"),
    ("Grupo E", 1, "Alemanha", "Curaçau", "14/06/2026", "14:00"),
    ("Grupo F", 1, "Holanda", "Japão", "14/06/2026", "17:00"),
    ("Grupo E", 1, "Costa do Marfim", "Equador", "14/06/2026", "20:00"),
    ("Grupo F", 1, "Suécia", "Tunísia", "14/06/2026", "23:00"),
    ("Grupo H", 1, "Espanha", "Cabo Verde", "15/06/2026", "13:00"),
    ("Grupo G", 1, "Bélgica", "Egito", "15/06/2026", "16:00"),
    ("Grupo H", 1, "Arábia Saudita", "Uruguai", "15/06/2026", "19:00"),
    ("Grupo G", 1, "Irã", "Nova Zelândia", "15/06/2026", "22:00"),
    ("Grupo I", 1, "França", "Senegal", "16/06/2026", "16:00"),
    ("Grupo I", 1, "Iraque", "Noruega", "16/06/2026", "19:00"),
    ("Grupo J", 1, "Argentina", "Argélia", "16/06/2026", "22:00"),
    ("Grupo J", 1, "Áustria", "Jordânia", "17/06/2026", "01:00"),
    ("Grupo K", 1, "Portugal", "RD Congo", "17/06/2026", "14:00"),
    ("Grupo L", 1, "Inglaterra", "Croácia", "17/06/2026", "17:00"),
    ("Grupo L", 1, "Gana", "Panamá", "17/06/2026", "20:00"),
    ("Grupo K", 1, "Uzbequistão", "Colômbia", "17/06/2026", "23:00"),
    # Rodada 2
    ("Grupo A", 2, "República Tcheca", "África do Sul", "18/06/2026", "13:00"),
    ("Grupo B", 2, "Suíça", "Bósnia e Herzegovina", "18/06/2026", "16:00"),
    ("Grupo B", 2, "Canadá", "Catar", "18/06/2026", "19:00"),
    ("Grupo A", 2, "México", "Coreia do Sul", "18/06/2026", "22:00"),
    ("Grupo D", 2, "EUA", "Austrália", "19/06/2026", "16:00"),
    ("Grupo C", 2, "Escócia", "Marrocos", "19/06/2026", "19:00"),
    ("Grupo C", 2, "Brasil", "Haiti", "19/06/2026", "21:30"),
    ("Grupo D", 2, "Turquia", "Paraguai", "20/06/2026", "00:30"),
    ("Grupo E", 2, "Alemanha", "Costa do Marfim", "20/06/2026", "14:00"),
    ("Grupo F", 2, "Holanda", "Suécia", "20/06/2026", "17:00"),
    ("Grupo E", 2, "Equador", "Curaçau", "20/06/2026", "21:00"),
    ("Grupo F", 2, "Tunísia", "Japão", "21/06/2026", "01:00"),
    ("Grupo H", 2, "Espanha", "Arábia Saudita", "21/06/2026", "13:00"),
    ("Grupo G", 2, "Bélgica", "Irã", "21/06/2026", "16:00"),
    ("Grupo H", 2, "Uruguai", "Cabo Verde", "21/06/2026", "19:00"),
    ("Grupo G", 2, "Nova Zelândia", "Egito", "21/06/2026", "22:00"),
    ("Grupo I", 2, "França", "Iraque", "22/06/2026", "16:00"),
    ("Grupo I", 2, "Noruega", "Senegal", "22/06/2026", "19:00"),
    ("Grupo J", 2, "Argentina", "Áustria", "22/06/2026", "22:00"),
    ("Grupo J", 2, "Jordânia", "Argélia", "23/06/2026", "01:00"),
    ("Grupo K", 2, "Portugal", "Uzbequistão", "23/06/2026", "13:00"),
    ("Grupo L", 2, "Inglaterra", "Gana", "23/06/2026", "16:00"),
    ("Grupo L", 2, "Panamá", "Croácia", "23/06/2026", "19:00"),
    ("Grupo K", 2, "Colômbia", "RD Congo", "23/06/2026", "22:00"),
    # Rodada 3
    ("Grupo B", 3, "Suíça", "Canadá", "24/06/2026", "16:00"),
    ("Grupo B", 3, "Bósnia e Herzegovina", "Catar", "24/06/2026", "16:00"),
    ("Grupo C", 3, "Escócia", "Brasil", "24/06/2026", "19:00"),
    ("Grupo C", 3, "Marrocos", "Haiti", "24/06/2026", "19:00"),
    ("Grupo A", 3, "República Tcheca", "México", "24/06/2026", "22:00"),
    ("Grupo A", 3, "África do Sul", "Coreia do Sul", "24/06/2026", "22:00"),
    ("Grupo E", 3, "Equador", "Alemanha", "25/06/2026", "13:00"),
    ("Grupo E", 3, "Curaçau", "Costa do Marfim", "25/06/2026", "13:00"),
    ("Grupo F", 3, "Japão", "Suécia", "25/06/2026", "16:00"),
    ("Grupo F", 3, "Tunísia", "Holanda", "25/06/2026", "16:00"),
    ("Grupo D", 3, "Austrália", "Paraguai", "25/06/2026", "19:00"),
    ("Grupo D", 3, "EUA", "Turquia", "25/06/2026", "19:00"),
    ("Grupo G", 3, "Nova Zelândia", "Bélgica", "26/06/2026", "13:00"),
    ("Grupo G", 3, "Egito", "Irã", "26/06/2026", "13:00"),
    ("Grupo I", 3, "Noruega", "França", "26/06/2026", "16:00"),
    ("Grupo I", 3, "Senegal", "Iraque", "26/06/2026", "16:00"),
    ("Grupo H", 3, "Cabo Verde", "Arábia Saudita", "26/06/2026", "21:00"),
    ("Grupo H", 3, "Uruguai", "Espanha", "26/06/2026", "21:00"),
    ("Grupo J", 3, "Áustria", "Argélia", "27/06/2026", "01:00"),
    ("Grupo J", 3, "Jordânia", "Argentina", "27/06/2026", "01:00"),
    ("Grupo L", 3, "Panamá", "Inglaterra", "27/06/2026", "16:00"),
    ("Grupo L", 3, "Croácia", "Gana", "27/06/2026", "16:00"),
    ("Grupo K", 3, "RD Congo", "Uzbequistão", "27/06/2026", "19:00"),
    ("Grupo K", 3, "Colômbia", "Portugal", "27/06/2026", "19:00"),
]

# Criar workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Copa do Mundo 2026"

# Estilos
header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center")

data_font = Font(name="Calibri", size=11)
data_alignment = Alignment(horizontal="center", vertical="center")

# Cores alternadas por grupo
group_colors = {
    "Grupo A": "DAEEF3",
    "Grupo B": "E2EFDA",
    "Grupo C": "FFF2CC",
    "Grupo D": "FCE4D6",
    "Grupo E": "D9E2F3",
    "Grupo F": "EDEDED",
    "Grupo G": "E2F0D9",
    "Grupo H": "FDE9D9",
    "Grupo I": "D6DCE4",
    "Grupo J": "F2DCDB",
    "Grupo K": "DAEEF3",
    "Grupo L": "EBF1DE",
}

thin_border = Border(
    left=Side(style="thin", color="999999"),
    right=Side(style="thin", color="999999"),
    top=Side(style="thin", color="999999"),
    bottom=Side(style="thin", color="999999"),
)

# Cabeçalhos
headers = ["Grupo", "Rodada", "Time Casa", "Time Visitante", "Data", "Hora"]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# Dados
for row_idx, jogo in enumerate(jogos, 2):
    grupo, rodada, casa, visitante, data, hora = jogo
    row_data = [grupo, f"{rodada}ª Rodada", casa, visitante, data, hora]

    color = group_colors.get(grupo, "FFFFFF")
    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = data_font
        cell.alignment = data_alignment
        cell.border = thin_border
        cell.fill = fill

# Ajustar largura das colunas
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 25
ws.column_dimensions["D"].width = 25
ws.column_dimensions["E"].width = 14
ws.column_dimensions["F"].width = 10

# Congelar painel (cabeçalho fixo)
ws.freeze_panes = "A2"

# Filtro automático
ws.auto_filter.ref = f"A1:F{len(jogos) + 1}"

# Salvar
output_path = r"c:\Users\Danilo\OneDrive\Documentos\2026\Python\BolaoCopa\Copa_do_Mundo_2026_Jogos.xlsx"
wb.save(output_path)
print(f"[OK] Arquivo salvo com sucesso em: {output_path}")
print(f"[INFO] Total de jogos: {len(jogos)}")
