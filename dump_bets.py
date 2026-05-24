import os
import time
import datetime
import requests
import schedule
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_ANON_KEY = os.getenv("SUPABASE_ANON_KEY")

if not SUPABASE_URL or not SUPABASE_ANON_KEY:
    print("[ERRO] Variáveis SUPABASE_URL e/ou SUPABASE_ANON_KEY não encontradas.")
    exit(1)

HEADERS = {
    "apikey": SUPABASE_ANON_KEY,
    "Authorization": f"Bearer {SUPABASE_ANON_KEY}",
    "Content-Type": "application/json"
}

def fetch_table(table_name):
    print(f"Buscando dados de {table_name}...")
    url = f"{SUPABASE_URL}/rest/v1/{table_name}?select=*"
    
    # Tentativa de paginação simples
    all_data = []
    offset = 0
    limit = 1000
    while True:
        paginated_url = f"{url}&limit={limit}&offset={offset}"
        response = requests.get(paginated_url, headers=HEADERS)
        if response.status_code != 200:
            print(f"[ERRO] Falha ao buscar {table_name}: {response.text}")
            break
        
        data = response.json()
        if not data:
            break
            
        all_data.extend(data)
        if len(data) < limit:
            break
        offset += limit
        
    return all_data

def generate_excel_dump():
    print(f"[{datetime.datetime.now()}] Iniciando dump de apostas...")
    
    # 1. Fetch data
    profiles = fetch_table("profiles")
    teams = fetch_table("teams")
    matches = fetch_table("matches")
    bets = fetch_table("bets")
    
    # 2. Build lookups
    profiles_lookup = {p['id']: p for p in profiles}
    teams_lookup = {t['id']: t for t in teams}
    matches_lookup = {m['id']: m for m in matches}
    
    # 3. Create Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Apostas - Dump"
    
    # Styles
    header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="Calibri", size=11)
    data_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="999999"), right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"), bottom=Side(style="thin", color="999999")
    )
    
    headers = [
        "ID Aposta", "Usuário", "Data Jogo", "Grupo", 
        "Time Casa", "Placar Casa (Aposta)", "Placar Visitante (Aposta)", "Time Visitante",
        "Pontos", "Data da Aposta"
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        
    # 4. Fill Data
    row_idx = 2
    for bet in bets:
        user = profiles_lookup.get(bet['user_id'], {}).get('username', 'Desconhecido')
        match = matches_lookup.get(bet['match_id'], {})
        
        home_team_id = match.get('home_team_id')
        away_team_id = match.get('away_team_id')
        home_team = teams_lookup.get(home_team_id, {}).get('name', 'Desconhecido')
        away_team = teams_lookup.get(away_team_id, {}).get('name', 'Desconhecido')
        group = match.get('group_letter', '-')
        match_date = match.get('match_date', '-')
        
        # Formatar datas se possivel
        if match_date != '-':
            try:
                # Exemplo: 2026-06-11T16:00:00+00:00 -> 11/06/2026 16:00
                dt = datetime.datetime.fromisoformat(match_date.replace('Z', '+00:00'))
                match_date = dt.strftime("%d/%m/%Y %H:%M")
            except:
                pass
                
        bet_date = bet.get('updated_at', bet.get('created_at', '-'))
        if bet_date != '-':
            try:
                dt = datetime.datetime.fromisoformat(bet_date.replace('Z', '+00:00'))
                bet_date = dt.strftime("%d/%m/%Y %H:%M:%S")
            except:
                pass
                
        row_data = [
            bet['id'],
            user,
            match_date,
            f"Grupo {group}",
            home_team,
            bet['home_score_bet'],
            bet['away_score_bet'],
            away_team,
            bet['points'],
            bet_date
        ]
        
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border
            
        row_idx += 1
        
    # Adjust column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 25
    ws.column_dimensions["I"].width = 10
    ws.column_dimensions["J"].width = 20
    
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{row_idx - 1}"
    
    # 5. Save File
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    os.makedirs("/app/backups", exist_ok=True)
    file_path = f"/app/backups/Dump_Apostas_{timestamp}.xlsx"
    wb.save(file_path)
    print(f"[OK] Dump concluído. Salvo em: {file_path}")

def main():
    print("Iniciando serviço de Dump de Apostas...")
    # Executa uma vez no início (para teste inicial)
    generate_excel_dump()
    
    # Agenda a execução em 3 horários: 12:00, 18:00 e 00:00
    schedule.every().day.at("12:00").do(generate_excel_dump)
    schedule.every().day.at("18:00").do(generate_excel_dump)
    schedule.every().day.at("00:00").do(generate_excel_dump)
    print("Jobs agendados para executar às 12:00, 18:00 e 00:00 todos os dias.")
    
    while True:
        schedule.run_pending()
        time.sleep(60)

if __name__ == "__main__":
    main()
